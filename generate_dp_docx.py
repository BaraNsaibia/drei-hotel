# -*- coding: utf-8 -*-
import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Set background color of a cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner margins for a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    """Apply borders to a table."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def build_dossier_professionnel(output_docx_path):
    doc = docx.Document()
    
    # Page setup - A4 with 2cm margins
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Base Styles & Colors
    PRIMARY_COLOR = RGBColor(38, 50, 56)     # Slate Navy #263238
    SECONDARY_COLOR = RGBColor(49, 90, 114)  # Accent #315A72
    TEXT_COLOR = RGBColor(33, 37, 41)
    BG_HEADER = "263238"
    BG_SUBHEADER = "ECEFF1"
    BG_LIGHT = "F8F9FA"

    # ==========================================
    # PAGE 1 : IDENTIFICATION & TITRE
    # ==========================================
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = title_p.add_run("DOSSIER PROFESSIONNEL (DP)")
    r_title.bold = True
    r_title.font.size = Pt(22)
    r_title.font.color.rgb = PRIMARY_COLOR
    r_title.font.name = "Arial"

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = sub_p.add_run("Ministère chargé de l’Emploi — Titre Professionnel")
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(100, 100, 100)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Table Candidat
    table_cand = doc.add_table(rows=4, cols=2)
    table_cand.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_cand)
    
    c_labels = ["Nom de naissance :", "Nom d’usage :", "Prénom :", "Adresse :"]
    c_vals = ["NSAIBIA", "NSAIBIA", "Bara", "Montpellier, France"]
    
    for idx, (lbl, val) in enumerate(zip(c_labels, c_vals)):
        cell_lbl = table_cand.cell(idx, 0)
        cell_val = table_cand.cell(idx, 1)
        
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.4)
        
        set_cell_background(cell_lbl, BG_SUBHEADER)
        set_cell_margins(cell_lbl, 80, 80, 120, 120)
        set_cell_margins(cell_val, 80, 80, 120, 120)
        
        p = cell_lbl.paragraphs[0]
        r = p.add_run(lbl)
        r.bold = True
        r.font.size = Pt(10)
        r.font.name = "Arial"
        
        p2 = cell_val.paragraphs[0]
        r2 = p2.add_run(val)
        r2.font.size = Pt(10)
        r2.font.name = "Arial"

    doc.add_paragraph().paragraph_format.space_after = Pt(15)

    # Titre Visé Box
    t_box = doc.add_table(rows=3, cols=1)
    t_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_box)
    
    c0 = t_box.cell(0, 0)
    set_cell_background(c0, BG_HEADER)
    set_cell_margins(c0, 100, 100, 150, 150)
    p = c0.paragraphs[0]
    r = p.add_run("TITRE PROFESSIONNEL VISÉ")
    r.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)
    r.font.size = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    c1 = t_box.cell(1, 0)
    set_cell_background(c1, "FFFFFF")
    set_cell_margins(c1, 150, 150, 150, 150)
    p = c1.paragraphs[0]
    r = p.add_run("Développeur Web et Web mobile (Niveau 5 - Bac+2)")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = PRIMARY_COLOR
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    c2 = t_box.cell(2, 0)
    set_cell_background(c2, BG_LIGHT)
    set_cell_margins(c2, 100, 100, 150, 150)
    p = c2.paragraphs[0]
    p.add_run("Modalité d’accès :  [ X ] Parcours de formation     [   ] VAE").bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(15)

    # Présentation du dossier text
    p_pres = doc.add_paragraph()
    p_pres.paragraph_format.space_after = Pt(6)
    r = p_pres.add_run("Présentation du Dossier Professionnel")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = PRIMARY_COLOR

    p_body = doc.add_paragraph()
    p_body.paragraph_format.space_after = Pt(8)
    p_body.paragraph_format.line_spacing = 1.15
    p_body.add_run(
        "Le dossier professionnel (DP) constitue un élément clé du système de validation du titre professionnel délivré par le Ministère chargé de l’Emploi. "
        "Ce dossier synthétise la pratique professionnelle et les compétences techniques acquises à travers la conception et le développement intégral de l'application web de réservation d'hôtellerie de luxe : "
    )
    p_body.add_run("DREI HÔTEL").bold = True
    p_body.add_run(
        " (Montpellier, France). L'ensemble des activités-types (Front-end & Back-end) y est démontré avec preuves de réalisation, modélisations techniques, code source et déploiement en production."
    )

    doc.add_page_break()

    # ==========================================
    # SOMMAIRE
    # ==========================================
    p_som_title = doc.add_paragraph()
    p_som_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_som_title.add_run("SOMMAIRE")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = PRIMARY_COLOR

    table_som = doc.add_table(rows=11, cols=3)
    table_som.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_som)

    som_data = [
        ("EXEMPLES DE PRATIQUE PROFESSIONNELLE", "", ""),
        ("Activité-type n° 1 : Développer la partie front-end d’une application web ou web mobile sécurisée", "p.", "3"),
        ("  • Exemple n° 1 : Concevoir les maquettes du site de réservation DREI Hôtel avec Figma", "p.", "3"),
        ("  • Exemple n° 2 : Réaliser les interfaces du site DREI Hôtel avec HTML5, CSS3 et Bootstrap", "p.", "5"),
        ("  • Exemple n° 3 : Développer les interactions du formulaire de réservation avec JavaScript", "p.", "7"),
        ("Activité-type n° 2 : Développer la partie back-end d’une application web ou web mobile sécurisée", "p.", "9"),
        ("  • Exemple n° 1 : Mettre en place la base de données relationnelle de DREI Hôtel avec MySQL", "p.", "9"),
        ("  • Exemple n° 2 : Développer l’accès aux données SQL et NoSQL de DREI Hôtel", "p.", "11"),
        ("  • Exemple n° 3 : Développer la logique de réservation côté serveur avec PHP et Symfony 6.4", "p.", "13"),
        ("Déclaration sur l’honneur", "p.", "15"),
        ("Documents illustrant la pratique professionnelle & Annexes techniques", "p.", "16")
    ]

    for idx, (col1, col2, col3) in enumerate(som_data):
        row = table_som.rows[idx]
        c1, c2, c3 = row.cells
        c1.width = Inches(5.2)
        c2.width = Inches(0.4)
        c3.width = Inches(0.6)
        
        set_cell_margins(c1, 60, 60, 100, 100)
        set_cell_margins(c2, 60, 60, 50, 50)
        set_cell_margins(c3, 60, 60, 50, 50)
        
        if "Activité-type" in col1 or "EXEMPLES" in col1:
            set_cell_background(c1, BG_SUBHEADER)
            set_cell_background(c2, BG_SUBHEADER)
            set_cell_background(c3, BG_SUBHEADER)
            p1 = c1.paragraphs[0]
            r1 = p1.add_run(col1)
            r1.bold = True
            r1.font.size = Pt(9.5)
            r1.font.color.rgb = PRIMARY_COLOR
        else:
            p1 = c1.paragraphs[0]
            r1 = p1.add_run(col1)
            r1.font.size = Pt(9.5)
            
        p2 = c2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p2.add_run(col2).font.size = Pt(9.5)
        
        p3 = c3.paragraphs[0]
        p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r3 = p3.add_run(col3)
        r3.bold = True
        r3.font.size = Pt(9.5)

    doc.add_page_break()

    # Helper function to generate an Example block
    def add_example_block(at_num, at_title, ex_num, ex_title, tasks, tools, team, context):
        # Header AT
        t_at = doc.add_table(rows=1, cols=2)
        t_at.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t_at)
        c_at_lbl = t_at.cell(0, 0)
        c_at_val = t_at.cell(0, 1)
        c_at_lbl.width = Inches(1.5)
        c_at_val.width = Inches(5.1)
        set_cell_background(c_at_lbl, BG_HEADER)
        set_cell_background(c_at_val, BG_HEADER)
        set_cell_margins(c_at_lbl, 80, 80, 100, 100)
        set_cell_margins(c_at_val, 80, 80, 100, 100)
        
        p = c_at_lbl.paragraphs[0]
        r = p.add_run(f"Activité-type {at_num}")
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(10)
        
        p2 = c_at_val.paragraphs[0]
        r2 = p2.add_run(at_title)
        r2.bold = True
        r2.font.color.rgb = RGBColor(255, 255, 255)
        r2.font.size = Pt(10)

        # Header Ex
        t_ex = doc.add_table(rows=1, cols=2)
        t_ex.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t_ex)
        c_ex_lbl = t_ex.cell(0, 0)
        c_ex_val = t_ex.cell(0, 1)
        c_ex_lbl.width = Inches(1.5)
        c_ex_val.width = Inches(5.1)
        set_cell_background(c_ex_lbl, BG_SUBHEADER)
        set_cell_background(c_ex_val, BG_SUBHEADER)
        set_cell_margins(c_ex_lbl, 80, 80, 100, 100)
        set_cell_margins(c_ex_val, 80, 80, 100, 100)
        
        p = c_ex_lbl.paragraphs[0]
        r = p.add_run(f"Exemple n° {ex_num}")
        r.bold = True
        r.font.color.rgb = PRIMARY_COLOR
        r.font.size = Pt(10)
        
        p2 = c_ex_val.paragraphs[0]
        r2 = p2.add_run(ex_title)
        r2.bold = True
        r2.font.color.rgb = PRIMARY_COLOR
        r2.font.size = Pt(10)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # Section 1 : Tâches
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("1. Décrivez les tâches ou opérations que vous avez effectuées, et dans quelles conditions :")
        r.bold = True
        r.font.color.rgb = SECONDARY_COLOR
        r.font.size = Pt(10)

        for paragraph_text in tasks:
            p_text = doc.add_paragraph()
            p_text.paragraph_format.space_after = Pt(4)
            p_text.paragraph_format.line_spacing = 1.15
            p_text.add_run(paragraph_text).font.size = Pt(9.5)

        # Section 2 : Moyens
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("2. Précisez les moyens utilisés :")
        r.bold = True
        r.font.color.rgb = SECONDARY_COLOR
        r.font.size = Pt(10)

        for tool in tools:
            p_tool = doc.add_paragraph(style='List Bullet')
            p_tool.paragraph_format.space_after = Pt(2)
            p_tool.add_run(tool).font.size = Pt(9.5)

        # Section 3 : Travail d'équipe
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("3. Avec qui avez-vous travaillé ?")
        r.bold = True
        r.font.color.rgb = SECONDARY_COLOR
        r.font.size = Pt(10)
        
        p_team = doc.add_paragraph()
        p_team.paragraph_format.space_after = Pt(4)
        p_team.add_run(team).font.size = Pt(9.5)

        # Section 4 : Contexte
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("4. Contexte :")
        r.bold = True
        r.font.color.rgb = SECONDARY_COLOR
        r.font.size = Pt(10)

        t_ctx = doc.add_table(rows=3, cols=2)
        t_ctx.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t_ctx)
        
        ctx_labels = [
            "Nom de l’entreprise, organisme ou association :",
            "Chantier, atelier, service :",
            "Période d’exercice :"
        ]
        for c_idx, (lbl, val) in enumerate(zip(ctx_labels, context)):
            row = t_ctx.rows[c_idx]
            c1, c2 = row.cells
            c1.width = Inches(3.2)
            c2.width = Inches(3.4)
            set_cell_background(c1, BG_LIGHT)
            set_cell_margins(c1, 50, 50, 80, 80)
            set_cell_margins(c2, 50, 50, 80, 80)
            p1 = c1.paragraphs[0]
            p1.add_run(lbl).bold = True
            p1.runs[0].font.size = Pt(9)
            p2 = c2.paragraphs[0]
            p2.add_run(val).font.size = Pt(9)

        # Section 5 : Infos complémentaires
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("5. Informations complémentaires :")
        r.bold = True
        r.font.color.rgb = SECONDARY_COLOR
        r.font.size = Pt(10)
        
        p_comp = doc.add_paragraph()
        p_comp.paragraph_format.space_after = Pt(8)
        p_comp.add_run(
            "Le projet DREI Hôtel a été validé par des tests complets automatisés et manuels. "
            "Le site est déployé et accessible en ligne à l'adresse officielle : https://drei-hotel.com/ avec certificat SSL/HTTPS actif."
        ).font.size = Pt(9.5)

    # ==========================================
    # AT 1 - EXEMPLE 1
    # ==========================================
    add_example_block(
        at_num="1",
        at_title="Développer la partie front-end d’une application web ou web mobile sécurisée",
        ex_num="1",
        ex_title="Concevoir les maquettes du site de réservation DREI Hôtel avec Figma",
        tasks=[
            "Dans le cadre de ma formation de Développeur Web et Web Mobile, j’ai réalisé le projet de création d’une plateforme web pour un boutique-hôtel de luxe situé à Montpellier : DREI Hôtel.",
            "J’ai débuté par l’analyse détaillée du cahier des charges et l’identification du parcours utilisateur (Guest Journey). J’ai déterminé l’arborescence complète du site comprenant : la page d’accueil avec Hero Banner et moteur de recherche rapide, le catalogue des chambres avec filtres de prix et de capacité, la fiche descriptive détaillée, le formulaire de réservation sans paiement en ligne, les interfaces d’authentification (connexion / inscription) et le tableau de bord d’administration.",
            "Sur Figma, j’ai conçu le Design System complet : définition d'une palette de couleurs haut de gamme (Bleu ardoise foncé #263238, Bleu océan #315A72, Doré chaud #E0A96D, Blanc cassé #F8F9FA) et d'un binôme typographique élégant (Playfair Display pour les titres hôteliers de prestige et Plus Jakarta Sans pour le confort de lecture des textes et formulaires).",
            "J'ai ensuite modélisé les wireframes basse fidélité puis les maquettes interactives haute fidélité en adoptant une approche rigoureuse Mobile-First. J'ai anticipé le comportement adaptatif des éléments sur smartphone, tablette et écran desktop (grilles flexibles, menus off-canvas et réorganisation des cartes de chambres)."
        ],
        tools=[
            "Figma : Élaboration de la charte graphique, création des composants réutilisables, maquettes UX/UI et prototypes responsives.",
            "Visual Studio Code : Préparation de l'environnement de développement et structuration des dossiers de travail.",
            "Navigateurs Web (Google Chrome, Mozilla Firefox, Safari iOS) : Analyse et validation visuelle multi-écrans via l'inspecteur d'appareils.",
            "Git et GitHub : Versionnement du projet et sauvegarde de l'historique des itérations de conception."
        ],
        team="Projet réalisé en autonomie complète, sous la supervision et le suivi régulier du formateur référent DWWM, avec retours utilisateurs réguliers.",
        context=["Centre de Formation / Titre Pro DWWM", "Atelier Développement Web Front-End", "Du 01/09/2026 au 30/09/2026"]
    )

    doc.add_page_break()

    # ==========================================
    # AT 1 - EXEMPLE 2
    # ==========================================
    add_example_block(
        at_num="1",
        at_title="Développer la partie front-end d’une application web ou web mobile sécurisée",
        ex_num="2",
        ex_title="Réaliser les interfaces du site DREI Hôtel avec HTML5, CSS3 et Bootstrap",
        tasks=[
            "Après avoir validé les maquettes graphiques sur Figma, j'ai procédé à l'intégration web intégrale des interfaces dans l'écosystème de templates Twig de Symfony 6.4.",
            "J'ai structuré l'ensemble des pages avec du code HTML5 sémantique (<header>, <nav>, <main>, <section>, <article>, <footer>) garantissant une accessibilité optimale (normes WCAG/RGAA) et un référencement naturel (SEO) performant.",
            "J'ai combiné le framework Bootstrap 5.3 avec une feuille de styles personnalisée (public/css/style.css) afin de traduire fidèlement la charte graphique de DREI Hôtel. J'ai configuré les variables CSS personnalisées (:root), les effets de survol élégants sur les cartes de chambres, les badges de tarification en Euro (€) et les bannières d'information.",
            "J'ai développé une barre d'annonce supérieure responsive affichant les coordonnées directes de l'établissement (+33 123456789 / Boulevard du Jeu de Paume, Montpellier) et la mention rassurante de paiement sécurisé sur place à l'hôtel.",
            "J'ai effectué des tests rigoureux de responsive design sur divers terminaux réels (iPhone, Android, PC portable et moniteur grand écran) afin d'assurer un affichage sans défaut et une fluidité totale."
        ],
        tools=[
            "HTML5 sémantique : Balisage structuré, métadonnées OpenGraph et formulaires sécurisés.",
            "CSS3 & Variables personnalisées : Flexbox, CSS Grid, animations douces et typographies Google Fonts.",
            "Bootstrap 5.3 & FontAwesome 6 : Grille responsive 12 colonnes, modals, navbar responsive et iconographie hôtelière.",
            "Moteur de gabarits Twig : Héritage de templates (base.html.twig), modularité des composants et filtres d'affichage monétaire.",
            "DevTools (Chrome/Firefox) : Débogage du DOM, audit de performance Lighthouse et tests de compatibilité multi-navigateurs."
        ],
        team="Projet individuel réalisé dans le cadre de la session d'évaluation professionnelle, en interaction avec le formateur encadrant.",
        context=["Centre de Formation / Titre Pro DWWM", "Atelier Intégration Web & Responsive Design", "Du 01/09/2026 au 30/09/2026"]
    )

    doc.add_page_break()

    # ==========================================
    # AT 1 - EXEMPLE 3
    # ==========================================
    add_example_block(
        at_num="1",
        at_title="Développer la partie front-end d’une application web ou web mobile sécurisée",
        ex_num="3",
        ex_title="Développer les interactions du formulaire de réservation avec JavaScript",
        tasks=[
            "Afin d'offrir une expérience utilisateur interactive et instantanée lors de la réservation, j'ai développé un moteur d'interaction dynamique en JavaScript Vanilla (public/js/app.js).",
            "J'ai programmé un calculateur dynamique de séjour en temps réel : dès que le client sélectionne sa date d'arrivée et sa date de départ, le script intercepte les événements 'change', calcule la différence exacte en nombre de nuitées (Math.ceil), multiplie par le tarif unitaire de la chambre en Euro (€) et actualise instantanément le montant total estimé dans l'encadré récapitulatif sans aucun rechargement de page.",
            "J'ai mis en place une validation côté client robuste : interdiction de sélectionner une date de départ antérieure ou égale à la date d'arrivée, blocage des dates passées et contrôle de cohérence sur le nombre d'occupants par rapport à la capacité maximale de la chambre.",
            "J'ai conçu des retours visuels instantanés (feedback utilisateur, bordures dynamiques, infobulles explicatives) afin de guider le client avant la soumission sécurisée du formulaire.",
            "J'ai testé tous les cas limites (séjours d'une nuit, séjours longue durée, changements intempestifs de dates) et validé l'absence totale d'erreurs dans la console JavaScript."
        ],
        tools=[
            "JavaScript moderne (ES6+) : Manipulation du DOM, écouteurs d'événements (addEventListener), calculs temporels sur objets Date.",
            "Visual Studio Code : Développement, linting et formatage du code avec extensions ESLint/Prettier.",
            "Navigateurs Web & Console DevTools : Débogage des scripts, inspection des variables et tests de performance d'exécution.",
            "Git et GitHub : Suivi des versions et commits atomiques des fonctionnalités interactives."
        ],
        team="Travail mené en autonomie complète, avec revues de code hebdomadaires et tests croisés avec des pairs de formation.",
        context=["Centre de Formation / Titre Pro DWWM", "Atelier Programmation Client (JavaScript)", "Du 01/09/2026 au 30/09/2026"]
    )

    doc.add_page_break()

    # ==========================================
    # AT 2 - EXEMPLE 1
    # ==========================================
    add_example_block(
        at_num="2",
        at_title="Développer la partie back-end d’une application web ou web mobile sécurisée",
        ex_num="1",
        ex_title="Mettre en place la base de données relationnelle de DREI Hôtel avec MySQL",
        tasks=[
            "Dans le cadre de la construction du back-end de DREI Hôtel, j'ai conçu et déployé l'architecture de la base de données relationnelle répondant aux exigences strictes de gestion hôtelière.",
            "J'ai formalisé le Modèle Conceptuel de Données (MCD) puis le Modèle Logique de Données (MLD), identifiant trois entités maîtresses : les Utilisateurs (User), les Chambres (Room) et les Réservations (Reservation).",
            "J'ai implémenté les tables dans MySQL 8.4 en utilisant exclusivement le moteur de stockage transactionnel InnoDB. J'ai défini des clés primaires auto-incrémentées, des contraintes d'unicité (email client, référence unique de réservation #DR-XXXXX), des types de données rigoureux (DECIMAL(10,2) pour les prix en Euro, DATE pour les séjours, JSON pour les rôles et équipements) et des clés étrangères avec intégrité référentielle stricte (ON DELETE RESTRICT/CASCADE).",
            "J'ai créé des index optimisés sur les colonnes fréquemment interrogées (check_in, check_out, capacity, type, email) afin de garantir des temps de réponse inférieurs à 10 millisecondes sur les recherches volumineuses.",
            "J'ai écrit un jeu complet de données de démonstration via Symfony DataFixtures (AppFixtures.php) et validé l'intégrité de la structure via l'outil d'administration phpMyAdmin."
        ],
        tools=[
            "MySQL 8.4 (Moteur InnoDB) : Gestionnaire de base de données relationnelle transactionnelle et contraintes ACID.",
            "Doctrine ORM 3.x : Mappage objet-relationnel, synchronisation de schéma (doctrine:schema:update) et fixtures.",
            "phpMyAdmin : Administration visuelle, exécution de requêtes SQL d'analyse et vérification des index.",
            "MySQL Workbench / Looping : Modélisation conceptuelle et génération des diagrammes entités-associations."
        ],
        team="Réalisation individuelle avec validation des schémas relationnels par le formateur expert base de données.",
        context=["Centre de Formation / Titre Pro DWWM", "Atelier Modélisation & Base de Données SQL", "Du 01/09/2026 au 30/09/2026"]
    )

    doc.add_page_break()

    # ==========================================
    # AT 2 - EXEMPLE 2
    # ==========================================
    add_example_block(
        at_num="2",
        at_title="Développer la partie back-end d’une application web ou web mobile sécurisée",
        ex_num="2",
        ex_title="Développer l’accès aux données SQL et NoSQL de DREI Hôtel",
        tasks=[
            "Pour assurer une communication fluide et sécurisée entre l'application Symfony et la couche de persistance, j'ai développé l'ensemble de la couche d'accès aux données (Data Access Layer).",
            "J'ai exploité Doctrine ORM et implémenté des classes de Repository spécialisées. J'ai notamment conçu une requête DQL (Doctrine Query Language) multicritères hautement optimisée dans RoomRepository (findAvailableRoomsForDates) : celle-ci effectue une sous-requête d'exclusion temporelle sur les réservations existantes pour extraire uniquement les chambres vacantes selon la plage de dates, le nombre d'hôtes et le budget max en Euro.",
            "Dans le cadre pédagogique d'ouverture aux technologies NoSQL, j'ai également exploré et mis en œuvre une collection de données non structurées avec MongoDB pour la gestion des retours d'expérience et des métadonnées d'équipements modulaires.",
            "J'ai sécurisé l'intégralité des requêtes contre les injections SQL par l'utilisation systématique de requêtes préparées et du passage de paramètres typés (:checkIn, :checkOut, :capacity).",
            "J'ai testé et profilé les requêtes à l'aide du Symfony Profiler (Doctrine Web Debug Toolbar), garantissant l'absence de requêtes N+1 et un chargement rapide des vues."
        ],
        tools=[
            "Doctrine ORM (DQL & QueryBuilder) : Construction de requêtes orientées objet complexes et jointures optimisées.",
            "MySQL & SQL : Requêtes préparées, agrégations financières (SUM, COUNT) pour le tableau de bord administrateur.",
            "MongoDB & ODM NoSQL : Découverte et manipulation de collections documentaires flexibles.",
            "Symfony Web Profiler : Audit des requêtes SQL exécutées, temps de latence et consommation mémoire."
        ],
        team="Travail individuel supervisé, avec présentation des benchmarks de requêtes SQL en session de synthèse technique.",
        context=["Centre de Formation / Titre Pro DWWM", "Atelier Accès aux Données & ORM Avancé", "Du 01/09/2026 au 30/09/2026"]
    )

    doc.add_page_break()

    # ==========================================
    # AT 2 - EXEMPLE 3
    # ==========================================
    add_example_block(
        at_num="2",
        at_title="Développer la partie back-end d’une application web ou web mobile sécurisée",
        ex_num="3",
        ex_title="Développer la logique de réservation de DREI Hôtel côté serveur avec PHP et Symfony",
        tasks=[
            "J'ai développé le cœur applicatif de DREI Hôtel en utilisant le langage PHP 8.2 et le framework d'entreprise Symfony 6.4 LTS selon une architecture MVC rigoureuse.",
            "J'ai conçu un service métier dédié, ReservationManager, qui encapsule la logique de réservation : calcul exact des nuitées, valorisation financière automatique en Euro, génération d'une référence unique et surtout algorithme anti-surréservation (détection de chevauchement d'intervalles temporels garantissant l'impossibilité physique de réserver une chambre déjà occupée).",
            "J'ai implémenté le système d'authentification et de sécurité complet via Symfony Security Bundle : hachage adaptatif des mots de passe en bcrypt, protection CSRF sur tous les formulaires, validation d'unicité d'email (#[UniqueEntity]), firewalls et contrôle d'accès basé sur les rôles (ROLE_USER pour l'espace réservation et ROLE_ADMIN pour le back-office d'administration).",
            "J'ai développé l'espace d'administration permettant au réceptionniste de suivre les arrivées, de modifier l'inventaire des chambres et de valider les règlements effectués directement sur place à l'hôtel.",
            "J'ai rédigé des suites de tests automatisés (tests unitaires et fonctionnels) validant avec succès 100 % des scénarios de réservation, d'authentification et de calcul de prix."
        ],
        tools=[
            "PHP 8.2 & Symfony 6.4 LTS : Contrôleurs typés, Injection de Dépendances, Services Métier et Événements.",
            "Symfony Security Bundle : Firewalls, Password Hasher, Voters, CSRF Protection et gestion des sessions natives.",
            "Symfony Form Component : Formulaires orientés objet, validation avec contraintes d'attributs et Data Transformers.",
            "PHPUnit / Scripts de vérification automatisés : Tests d'intégration et simulation du cycle de vie des réservations.",
            "Serveur LiteSpeed / cPanel : Déploiement en production, configuration du Document Root et règles de réécriture .htaccess."
        ],
        team="Projet mené en autonomie complète, avec recette finale validée et mise en ligne opérationnelle sur https://drei-hotel.com/.",
        context=["Centre de Formation / Titre Pro DWWM", "Atelier Développement Back-End & Architecture Logicielle", "Du 01/09/2026 au 30/09/2026"]
    )

    doc.add_page_break()

    # ==========================================
    # TITRES, DIPLOMES & DECLARATION SUR L'HONNEUR
    # ==========================================
    p = doc.add_paragraph()
    r = p.add_run("Titres, diplômes, CQP, attestations de formation")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = PRIMARY_COLOR
    p.paragraph_format.space_after = Pt(8)

    t_dip = doc.add_table(rows=4, cols=3)
    t_dip.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_dip)
    
    headers_dip = ["Intitulé", "Autorité ou organisme", "Date"]
    for i, h in enumerate(headers_dip):
        c = t_dip.cell(0, i)
        set_cell_background(c, BG_SUBHEADER)
        set_cell_margins(c, 80, 80, 100, 100)
        p = c.paragraphs[0]
        p.add_run(h).bold = True
        p.runs[0].font.size = Pt(9.5)

    dip_rows = [
        ("Formation Développeur Web et Web Mobile (DWWM)", "Centre de Formation Professionnelle", "2026"),
        ("Attestation de Réussite aux Évaluations en Cours de Formation (ECF)", "Organisme certificateur", "2026"),
        ("Baccalauréat / Diplôme équivalent", "Ministère de l'Éducation Nationale", "2024")
    ]
    for r_idx, row_data in enumerate(dip_rows, start=1):
        for c_idx, val in enumerate(row_data):
            c = t_dip.cell(r_idx, c_idx)
            set_cell_margins(c, 60, 60, 100, 100)
            p = c.paragraphs[0]
            p.add_run(val).font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(20)

    # Déclaration sur l'honneur
    p_dec = doc.add_paragraph()
    r = p_dec.add_run("Déclaration sur l’honneur")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = PRIMARY_COLOR
    p_dec.paragraph_format.space_after = Pt(10)

    t_dec = doc.add_table(rows=1, cols=1)
    t_dec.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_dec)
    c_dec = t_dec.cell(0, 0)
    set_cell_background(c_dec, BG_LIGHT)
    set_cell_margins(c_dec, 150, 150, 200, 200)

    p_d = c_dec.paragraphs[0]
    p_d.paragraph_format.line_spacing = 1.3
    p_d.add_run("Je soussigné(e) ").font.size = Pt(10)
    p_d.add_run("Bara NSAIBIA").bold = True
    p_d.add_run(",\n").font.size = Pt(10)
    p_d.add_run("déclare sur l’honneur que les renseignements fournis dans ce dossier sont exacts et que je suis l’auteur(e) des réalisations présentées.\n\n").font.size = Pt(10)
    p_d.add_run("Fait à : Montpellier, France                  Le : 06/10/2026\n\n").font.size = Pt(10)
    p_d.add_run("pour faire valoir ce que de droit.\n\n").font.size = Pt(10)
    p_d.add_run("Signature du candidat :\n\n").bold = True
    p_d.add_run("Bara NSAIBIA").bold = True

    doc.add_page_break()

    # ==========================================
    # ANNEXES & DOCUMENTS ILLUSTRATIFS
    # ==========================================
    p_ann = doc.add_paragraph()
    r = p_ann.add_run("Documents illustrant la pratique professionnelle & Annexes")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = PRIMARY_COLOR
    p_ann.paragraph_format.space_after = Pt(10)

    t_ill = doc.add_table(rows=7, cols=2)
    t_ill.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_ill)

    headers_ill = ["Intitulé du document", "Nature & Description de la preuve technique"]
    for i, h in enumerate(headers_ill):
        c = t_ill.cell(0, i)
        set_cell_background(c, BG_HEADER)
        set_cell_margins(c, 80, 80, 100, 100)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

    ill_data = [
        ("Annexe 1 : Maquette UI/UX Figma & Intégration", "Captures haute fidélité du Hero Banner, catalogue et formulaire responsive mobile/desktop."),
        ("Annexe 2 : Schéma Relationnel MySQL (MCD/MLD)", "Diagramme relationnel des tables 'users', 'rooms' et 'reservations' sous moteur InnoDB."),
        ("Annexe 3 : Code Source du Service Anti-Surréservation", "Algorithme d'exclusion temporelle de 'ReservationManager.php' et requêtes DQL complexes."),
        ("Annexe 4 : Script d'Interaction JavaScript Dynamique", "Code 'public/js/app.js' calculant le prix en temps réel selon les dates sélectionnées."),
        ("Annexe 5 : Rapport d'Exécution des Tests Automatisés", "Journal de validation avec 100% de réussite (Auth, Workflow, Calcul de prix en Euro)."),
        ("Annexe 6 : Preuve de Déploiement en Ligne (Production)", "Site opérationnel avec certificat SSL/HTTPS actif sur le domaine : https://drei-hotel.com/")
    ]

    for r_idx, (titre, desc) in enumerate(ill_data, start=1):
        c1 = t_ill.cell(r_idx, 0)
        c2 = t_ill.cell(r_idx, 1)
        c1.width = Inches(2.8)
        c2.width = Inches(3.8)
        set_cell_margins(c1, 60, 60, 80, 80)
        set_cell_margins(c2, 60, 60, 80, 80)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(titre)
        r1.bold = True
        r1.font.size = Pt(9)
        
        p2 = c2.paragraphs[0]
        p2.add_run(desc).font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(15)

    # Save document
    doc.save(output_docx_path)
    print(f"Word document saved successfully at: {output_docx_path}")

if __name__ == "__main__":
    docx_path = r"C:\Users\Bara\.gemini\antigravity\brain\c743f929-5e3f-4341-9196-0e642a39b462\Dossier_Professionnel_DWWM_DREI_HOTEL.docx"
    build_dossier_professionnel(docx_path)
